import asyncio
import hashlib
import hmac
import io
import json
import logging
import secrets
import time
import zipfile
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, Header, UploadFile, File, Form, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import Response, FileResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy import select, delete
from PIL import Image
from pypdf import PdfReader
from docx import Document
from .config import settings
from .db import DB, init_db, Session, Question, Answer, Observation, Snapshot, Feelings, Consent, Job, Report, uid
from .schemas import SessionInput, TranscriptInput, FeelingsInput, Heartbeat, ObservationInput, ConsentInput, ExplanationInput, ViolationInput, MCQSubmission
from .ai import extract_profile, sanitize, generate_questions, aggregate
from .storage import storage
from .monitoring import validation_status
from .proctoring import load_config, describe, summarize
from .jobs import worker
from .accounts import actor, require_user, hashed, organization_recruiters, router as auth_router
from .db import User, LoginSession, Assignment, Examination
from .portal import router as portal_router
from .recruitment import router as recruitment_router
from .candidate_workspace import router as candidate_router
from .applications import router as applications_router
from .exam_policy import apply_policy, release_lock, notify, system_event, device_warning

logger = logging.getLogger('interview')


def digest(token):
    return hashlib.sha256(token.encode()).hexdigest()


def report_token(sid, role):
    return hmac.new(settings.app_secret.encode(), f'{sid}:{role}'.encode(), hashlib.sha256).hexdigest()

def session_token(sid):
    return hmac.new(settings.app_secret.encode(), f'account-session:{sid}'.encode(), hashlib.sha256).hexdigest()


def cleanup():
    with DB() as db:
        db.execute(delete(LoginSession).where(LoginSession.expires<=time.time()))
        db.commit()
        for s in db.scalars(select(Session)).all():
            if s.expires <= time.time() or s.status == 'deleting':
                try:
                    purge(db, s)
                except Exception:
                    s.status = 'deleting'
                    db.commit()
        expires_before=time.time()-settings.evidence_retention_days*86400
        for snap in db.scalars(select(Snapshot)).all():
            if snap.body.get('captured_at',0) and snap.body['captured_at']<=expires_before:
                try:
                    storage.delete(snap.body['key']);db.delete(snap)
                except Exception:
                    logger.warning('Could not remove expired evidence snapshot %s',snap.id)
        db.commit()
        active = {j.body.get('path') for j in db.scalars(select(Job)).all() if j.body.get('status') in ['queued', 'running']}
    for p in Path(settings.data_dir).glob('audio-*'):
        if str(p) not in active and time.time() - p.stat().st_mtime > 3600:
            p.unlink(missing_ok=True)


def purge(db, session):
    session.status = 'deleting'
    db.commit()  # Revoke access even if remote deletion must be retried.
    for snap in db.scalars(select(Snapshot).where(Snapshot.session_id == session.id)).all():
        storage.delete(snap.body['key'])
    for job in db.scalars(select(Job).where(Job.session_id == session.id)).all():
        if job.body.get('path'):
            Path(job.body['path']).unlink(missing_ok=True)
    for assignment in db.scalars(select(Assignment).where(Assignment.session_id==session.id)):
        assignment.body=assignment.body|{'status':'attempt_deleted'}
        assignment.session_id=None
    from .db import ApplicationRound
    for round_entry in db.scalars(select(ApplicationRound).where(ApplicationRound.interview_session_id==session.id)):
        round_entry.interview_session_id=None
    db.flush()
    db.delete(session)
    db.commit()


async def janitor(stop):
    while not stop.is_set():
        await asyncio.to_thread(cleanup)
        try:
            await asyncio.wait_for(stop.wait(), timeout=300)
        except asyncio.TimeoutError:
            pass


@asynccontextmanager
async def lifespan(app):
    init_db()
    stop = asyncio.Event()
    from .email_delivery import email_worker
    tasks = [asyncio.create_task(worker(stop)), asyncio.create_task(janitor(stop)),asyncio.create_task(email_worker(stop))]
    yield
    stop.set()
    for task in tasks:
        task.cancel()
    await asyncio.gather(*tasks, return_exceptions=True)


app = FastAPI(title='Interview Studio', lifespan=lifespan)
app.include_router(auth_router)
app.include_router(portal_router)
app.include_router(recruitment_router)
app.include_router(candidate_router)
app.include_router(applications_router)
from .screening import router as screening_router
app.include_router(screening_router)
from .assessment_history import effective_evaluation

from .rounds import router as rounds_router
app.include_router(rounds_router)


rate_windows = defaultdict(deque)


@app.middleware('http')
async def security(request: Request, call_next):
    start = time.monotonic()
    if request.method in ['POST', 'PUT', 'PATCH', 'DELETE']:
        try:
            if int(request.headers.get('content-length', '0')) > 13*1024*1024:
                return Response('Request too large', status_code=413)
        except ValueError:
            return Response('Invalid content length', status_code=400)
        # IP plus broad operation group; intentionally excludes report/session tokens from logs.
        path = request.url.path
        group = 'heartbeat' if path.endswith('heartbeat') else 'parse' if path.endswith('/parse') else 'ai' if path.endswith(('/evaluate', '/questions')) else 'create' if path == '/api/sessions' else 'write'
        key = (request.client.host if request.client else 'unknown', group)
        window = rate_windows[key]
        now = time.time()
        while window and window[0] < now-60:
            window.popleft()
        limits = {'heartbeat': 120, 'parse': 10, 'ai': 20, 'create': 10, 'write': 90, 'login':10}
        if path.startswith('/api/auth/'):
            group='login';key=(request.client.host if request.client else 'unknown',group);window=rate_windows[key]
            while window and window[0]<now-60:window.popleft()
        if len(window) >= limits[group]:
            return Response('Too many requests. Try again shortly.', status_code=429)
        window.append(now)
    response = await call_next(request)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Permissions-Policy'] = 'camera=(self), microphone=(self)'
    logger.info(json.dumps({'method': request.method, 'status': response.status_code, 'elapsed_ms': round((time.monotonic()-start)*1000)}))
    return response


@app.middleware('http')
async def authentication(request:Request,call_next):
    user=None;login=None
    cookie=request.cookies.get('interview_auth')
    with DB() as db:
        if cookie:
            login=db.get(LoginSession,hashed(cookie))
            if login and login.expires>time.time():user=db.get(User,login.user_id)
        if request.method not in ['GET','HEAD','OPTIONS']:
            origin=request.headers.get('origin')
            if origin and origin.rstrip('/')!=str(request.base_url).rstrip('/'):
                return Response('Cross-origin write rejected',status_code=403)
            if user and not request.url.path.startswith('/api/auth/'):
                if not hmac.compare_digest(login.csrf,hashed(request.headers.get('x-csrf-token',''))):
                    return Response('CSRF validation failed',status_code=403)
            if request.url.path=='/api/auth/logout' and login:
                if not hmac.compare_digest(login.csrf,hashed(request.headers.get('x-csrf-token',''))):return Response('CSRF validation failed',status_code=403)
                db.delete(login);db.commit()
    context=actor.set(user)
    try:return await call_next(request)
    finally:actor.reset(context)


def database():
    with DB() as db:
        yield db


def active_session(sid, db, authorization, mutable=True, evidence_only=False):
    s = db.get(Session, sid, with_for_update=mutable)
    token = authorization.removeprefix('Bearer ') if authorization else ''
    if not s or not hmac.compare_digest(s.token_hash, digest(token)):
        raise HTTPException(404, 'Session not found.')
    if s.body.get('user_id'):
        user=actor.get() if evidence_only else require_user('candidate')
        if not user or user.role!='candidate' or s.body['user_id']!=user.id:raise HTTPException(404,'Session not found.')
    if s.expires <= time.time() or s.status == 'deleting':
        raise HTTPException(410, 'Session expired or deleted.')
    if mutable and s.status in ['complete', 'terminated']:
        raise HTTPException(409, 'Assessment is finalized.')
    return s


def owned(cls, rid, db, authorization, mutable=True):
    row = db.get(cls, rid)
    if not row:
        raise HTTPException(404, 'Record not found.')
    return row, active_session(row.session_id, db, authorization, mutable)


def monitored(s, starting=False):
    if s.mode=="hr" and s.body.get("exam_id") and not starting and not s.body.get("exam_started_at"):
        raise HTTPException(409,"Complete the pre-exam verification first.")
    if s.body.get('security',{}).get('lock_until',0) and s.body['security']['lock_until']>time.time():
        raise HTTPException(423,'Examination temporarily locked. Wait for the server countdown to end.')
    if s.mode == 'hr':
        health = s.body.get('monitoring', {})
        if not health.get('healthy') or time.time()-health.get('received', 0) > 5:
            raise HTTPException(409, 'Working camera and device monitoring are required.')


def authorized_report(token, db, candidate=False):
    row = db.scalar(select(Report).where(Report.token_hash == digest(token)))
    if not row:
        raise HTTPException(404, 'Report not found.')
    s = db.get(Session, row.session_id)
    if not s or s.expires <= time.time() or s.status == 'deleting':
        raise HTTPException(410, 'Report expired or deleted.')
    if s.body.get('user_id'):
        user=require_user()
        allowed=(row.role=='candidate' and user.id==s.body['user_id']) or (row.role=='hr' and user.role in ('hr','admin') and s.body.get('hr_id') in organization_recruiters(user))
        if not allowed:raise HTTPException(403,'This report is not assigned to your account.')
    if candidate and row.role != 'candidate':
        raise HTTPException(403, 'Candidate access required.')
    return row, s


@app.get('/api/health')
def health():
    return {'status': 'ok'}


@app.get('/api/ready')
def ready(db=Depends(database)):
    db.execute(select(1))
    return {'status': 'ready', 'environment': settings.environment, 'ai_configured': bool(settings.gemini_api_key),
            'storage': 'supabase' if storage.remote else 'local development', 'detector': validation_status()}


@app.post('/api/resumes/parse')
async def parse_resume(file: UploadFile = File(...), name: str = Form('')):
    content = await file.read(5*1024*1024+1)
    if len(content) > 5*1024*1024:
        raise HTTPException(413, 'Resume must be under 5 MB.')
    ext = Path(file.filename or '').suffix.lower()
    try:
        if ext == '.pdf' and content.startswith(b'%PDF-'):
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 10:
                raise ValueError('Encrypted PDF or more than ten pages.')
            text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        elif ext == '.docx' and content.startswith(b'PK'):
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 20*1024*1024 or 'word/document.xml' not in z.namelist():
                    raise ValueError('Invalid or excessively expanded DOCX.')
            doc = Document(io.BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        else:
            raise ValueError('Upload a genuine PDF or DOCX file.')
        if len(text.strip()) < 30:
            raise ValueError('No usable text found. Scanned PDFs require OCR and are not supported.')
        if len(text) > 50000:
            raise ValueError('Resume text exceeds 50,000 characters.')
        return {'profile': extract_profile(text, name), 'notice': 'Review all entries and remove remaining identifying details. Original file and extracted text are not retained.'}
    except Exception as e:
        raise HTTPException(422, str(e) if isinstance(e, ValueError) else 'Could not read this document.')


@app.post('/api/sessions')
def create_session(data: SessionInput, db=Depends(database)):
    user=require_user('candidate')
    assignment=None;exam=None
    if data.mode=='hr':
        assignment=db.get(Assignment,data.assignment_id,with_for_update=True) if data.assignment_id else None
        if not assignment or assignment.candidate_id!=user.id:raise HTTPException(403,'HR screening requires an assigned examination.')
        if assignment.session_id:raise HTTPException(409,'An attempt already exists. Resume it from your dashboard.')
        if assignment.body.get('status')=='attempt_deleted':raise HTTPException(409,'Ask HR to assign a new examination.')

        exam=db.get(Examination,assignment.exam_id)
        if assignment.body.get('application_round_id'):
            from .db import ApplicationRound, Application
            round_entry=db.get(ApplicationRound,assignment.body['application_round_id'])
            application=db.get(Application,round_entry.application_id)
            if round_entry.status!='invited' or application.status=='rejected':
                raise HTTPException(409,'This round is not available to start.')
        from .accounts import organization_scope
        restriction=user.body.get('organization_restrictions',{}).get(organization_scope(db.get(User,exam.hr_id)))
        if restriction and (restriction['until'] is None or restriction['until']>time.time()):
            raise HTTPException(403,'Interview access is restricted for this company. Contact its HR team.')
        now=time.time(); start=exam.body.get('window_start_at',exam.body.get('opens_at',0)); end=exam.body.get('window_end_at',start+7200 if start else 0)
        if now<start:raise HTTPException(409,'This examination has not opened yet.')
        if end and now>end:
            assignment.body=assignment.body|{'attempt_used':True,'status':'expired','expired_at':now};db.commit()
            raise HTTPException(409,'This interview window has closed.')
        data.role=exam.body['role'];data.level=exam.body['level']
    if not data.processing_consent or data.mode == 'hr' and not data.evidence_consent:
        raise HTTPException(422, 'Required processing consent was not provided.')
    sid=uid();token = session_token(sid)
    profile = data.profile.model_dump()
    profile = {k: [sanitize(x, data.name) for x in v] if isinstance(v, list) else sanitize(v, data.name) for k, v in profile.items()}
    s = Session(id=sid,token_hash=digest(token), mode=data.mode, expires=time.time()+7*86400,
        body={'name': data.name, 'role': sanitize(data.role, data.name), 'level': data.level, 'profile': profile,
              'share_feelings': False, 'processing_consent': True, 'evidence_consent': data.evidence_consent,
              'remote_access_attestation': data.remote_access_attestation,
              'consent_version': '2026-09-proctor-v2', 'consented_at': time.time(), 'skill_focus': data.skill_focus,
              'proctor_config': exam.body['policy'] if exam else load_config(settings.proctor_config_file),
              'user_id':user.id,'hr_id':exam.hr_id if exam else None,'exam_id':exam.id if exam else None})
    db.add(s)
    db.flush()
    if assignment:
        if data.role.strip() != exam.body.get('role','').strip():
            raise HTTPException(409, 'The scheduled interview role cannot be changed.')
        assignment.session_id=s.id
        if assignment.body.get('application_round_id'):
            round_entry.interview_session_id=s.id;round_entry.status='in_progress'
            s.body=s.body|{'application_round_id':round_entry.id,'round_type':exam.body.get('round_type'),'evaluation_mode':exam.body.get('evaluation_mode','ai_conducted')}
        s.body=s.body|{'application_id':assignment.body.get('application_id'),'position_id':assignment.body.get('position_id')}
        db.add(Report(session_id=s.id,role='hr',token_hash=digest(report_token(s.id,'hr')),body={'live':True}))
    db.commit()
    return {'id': s.id, 'token': token}


@app.get('/api/sessions/{sid}')
def session_state(sid: str, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization, False)
    questions = sorted(db.scalars(select(Question).where(Question.session_id == sid)).all(), key=lambda q: q.body['sequence'])
    answers = db.scalars(select(Answer).where(Answer.session_id == sid)).all()
    return {'id': sid, 'mode': s.mode, 'round_type': s.body.get('round_type'), 'status': s.status, 'role': s.body['role'], 'name': s.body['name'],
            'questions': [{'id': q.id, **{k: v for k, v in q.body.items() if k != 'expected_points'}} for q in questions],
            'answers': [{'id': a.id, **a.body} for a in answers],
            'jobs': [{'id': j.id, **{k: v for k, v in j.body.items() if k != 'path'}} for j in db.scalars(select(Job).where(Job.session_id == sid))]}


@app.delete('/api/sessions/{sid}')
def cancel_session(sid: str, authorization: str = Header(''), db=Depends(database)):
    user = require_user('hr', 'admin')
    s = db.get(Session, sid)
    if not s or s.body.get('hr_id') not in organization_recruiters(user):
        raise HTTPException(404, 'Interview not found.')
    purge(db, s)
    return {'deleted': True}


@app.post('/api/sessions/{sid}/questions')
async def questions(sid: str, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization)
    existing = db.scalars(select(Question).where(Question.session_id == sid)).all()
    if not existing:
        from .round_questions import select_questions
        rows = select_questions(db,s.body.get('round_type'),s.body['level'])
        if rows is None:rows = await generate_questions(s.body['profile'], s.body['role'], s.body['level'], s.body.get('round_type'))
        # Deterministic primary keys protect concurrent retries.
        for i, q in enumerate(rows):
            db.merge(Question(id=f'{sid}-q{i+1}', session_id=sid, body=q | {'sequence': i+1, 'difficulty': s.body['level']}))
        db.commit()
    return session_state(sid, authorization, db)


@app.post('/api/sessions/{sid}/monitoring/heartbeat')
def heartbeat(sid: str, data: Heartbeat, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization)
    config = s.body.get('proctor_config') or load_config(settings.proctor_config_file)
    device_budget=config.get('device_result_max_age_ms',30000)
    device_enabled = validation_status()['passed']
    microphone_required = s.body.get('round_type') not in ('aptitude','technical')
    healthy = (data.camera and (data.microphone or not microphone_required) and data.face_age_ms < 5000 and data.evidence_ready
               and (not device_enabled or data.object_age_ms < device_budget and data.object_latency_ms <= device_budget)
               and (s.mode != 'hr' or data.page_visible and data.fullscreen and bool(s.body.get('storage_verified'))))
    prior = s.body.get('monitoring', {})
    if s.mode == 'hr' and (not healthy or prior and time.time()-prior.get('received', 0) > 5) and prior.get('healthy'):
        db.add(Observation(session_id=sid, body={'type': 'tracking_unavailable', 'timestamp': time.time()-s.created,
            'duration': 0, 'stage': 'monitoring', 'limitation': 'Monitoring interrupted; interview paused.'}))
    s.body = s.body | {'monitoring': data.model_dump() | {'healthy': healthy, 'received': time.time()}}
    s.status = 'active' if healthy else 'paused'
    db.commit()
    return {'healthy': healthy}


@app.get('/api/sessions/{sid}/monitoring')
def monitoring_summary(sid: str, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization, False)
    release_lock(db,s)
    config = s.body.get('proctor_config') or load_config(settings.proctor_config_file)
    rows = [{'id':o.id, **o.body} for o in db.scalars(select(Observation).where(Observation.session_id == sid))]
    config = {**config, 'device_detection_enabled': True, 'device_warning_only': not validation_status()['passed'],
              'microphone_required': s.body.get('round_type') not in ('aptitude','technical')}
    return {'config':config, 'summary':s.body.get('proctor_summary') or summarize(rows,config), 'events':[r for r in rows if r['type']!='gaze_sample'],
            'status':s.status, 'created':s.created, 'health':s.body.get('monitoring', {}),'security':s.body.get('security',{}),'server_now':time.time()}


@app.post('/api/sessions/{sid}/start')
def start_exam(sid:str, data:dict, authorization:str=Header(''),db=Depends(database)):
    s=active_session(sid,db,authorization);monitored(s, starting=True)
    if not data.get('rules_acknowledged'):raise HTTPException(422,'Acknowledge the examination rules.')
    if s.mode=='hr' and s.body.get('monitoring',{}).get('face_count')!=1:raise HTTPException(409,'Exactly one visible face is required to start.')
    if not s.body.get('exam_started_at'):
        s.body=s.body|{'exam_started_at':time.time(),'rules_acknowledged_at':time.time()};system_event(db,s,'exam_started');db.commit()
    return {'started':True}


@app.post('/api/sessions/{sid}/monitoring/storage-check')
def storage_check(sid: str, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization)
    key = f'{sid}/probe.jpg'
    out = io.BytesIO()
    Image.new('RGB', (1, 1)).save(out, 'JPEG')
    try:
        storage.put(key, out.getvalue())
        storage.get(key)
        storage.delete(key)
        s.body = s.body | {'storage_verified': time.time()}
        db.commit()
        return {'ready': True}
    except Exception:
        raise HTTPException(503, 'Evidence storage unavailable.')


@app.post('/api/sessions/{sid}/answers')
async def submit_answer(sid: str, question_id: str = Form(...), request_id: str = Form(...), transcript: str = Form(''),
                        duration: float = Form(0), audio: UploadFile | None = File(None),
                        authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization)
    monitored(s)
    q = db.get(Question, question_id)
    if not q or q.session_id != sid:
        raise HTTPException(404, 'Question not found.')
    earlier = [x for x in db.scalars(select(Question).where(Question.session_id == sid)) if x.body['sequence'] < q.body['sequence']]
    if any(not (prior_answer := db.get(Answer, f'{x.id}-answer')) or not prior_answer.body['approved']
           or prior_answer.body['evaluation']['status'] == 'pending' for x in earlier):
        raise HTTPException(409, 'Finish the previous question first.')
    limit=q.body.get('time_limit_seconds',90)
    if not 0 <= duration <= limit+1 or len(transcript) > 15000:
        raise HTTPException(422, 'Answer duration or length exceeds limit.')
    aid = f'{question_id}-answer'
    prior = db.get(Answer, aid)
    if prior:
        return {'id': prior.id, **prior.body}
    a = Answer(id=aid, session_id=sid, body={'question_id': question_id, 'request_id': request_id[:100],
        'transcript': transcript, 'duration': duration, 'approved': False, 'version': 1, 'metrics': {},
        'evaluation': {'status': 'pending'}, 'input': 'audio' if audio else 'typed'})
    db.add(a)
    if audio:
        raw = await audio.read(12*1024*1024+1)
        if len(raw) > 12*1024*1024:
            raise HTTPException(413, 'Recording exceeds 12 MB.')
        path = str(Path(settings.data_dir) / f'audio-{uid()}')
        Path(path).write_bytes(raw)
        job = Job(session_id=sid, body={'kind': 'transcribe', 'answer_id': aid, 'path': path, 'status': 'queued', 'version': 1})
        db.add(job)
    db.commit()
    return {'id': a.id, **a.body}


@app.get('/api/jobs/{jid}')
def get_job(jid: str, authorization: str = Header(''), db=Depends(database)):
    job, _ = owned(Job, jid, db, authorization, False)
    a = db.get(Answer, job.body['answer_id'])
    return {'id': jid, **{k: v for k, v in job.body.items() if k != 'path'}, 'answer': {'id': a.id, **a.body} if a else None}


@app.post('/api/sessions/{sid}/mcq')
def submit_mcq(sid:str, data:MCQSubmission, authorization:str=Header(''), db=Depends(database)):
    from .round_questions import grade_mcq
    s=active_session(sid,db,authorization)
    s=db.scalar(select(Session).where(Session.id==sid).with_for_update())
    monitored(s)
    q=db.get(Question,data.question_id)
    if not q or q.session_id!=sid:raise HTTPException(404,'Question not found.')
    options={o['id'] for o in q.body.get('options',[])}
    chosen=set(data.selected)
    if q.body.get('source')!='bank_mcq' or not options or not chosen<=options or (not chosen and not data.timed_out) or (chosen and not q.body.get('multiple_select') and len(chosen)!=1):
        raise HTTPException(422,'Choose a valid option for this question.')
    if data.duration>q.body.get('time_limit_seconds',90)+1:raise HTTPException(422,'Answer duration exceeds the question limit.')
    if data.timed_out and data.duration<q.body.get('time_limit_seconds',90):raise HTTPException(422,'Question time has not elapsed.')
    aid=q.id+'-answer'
    prior=db.get(Answer,aid)
    if prior and prior.body.get('approved') and prior.body.get('evaluation',{}).get('status')=='scored':
        return {'id':prior.id,**prior.body}
    earlier=list(db.scalars(select(Question).where(Question.session_id==sid)))
    answers={a.body.get('question_id'):a for a in db.scalars(select(Answer).where(Answer.session_id==sid))}
    if any(x.body['sequence']<q.body['sequence'] and (x.id not in answers or not answers[x.id].body.get('approved') or answers[x.id].body.get('evaluation',{}).get('status')=='pending') for x in earlier):
        raise HTTPException(409,'Finish the previous question first.')
    transcript=','.join(sorted(chosen))
    body={'question_id':q.id,'transcript':transcript,'duration':data.duration,'approved':True,
          'version':(prior.body.get('version',0) if prior else 0)+1,'metrics':{},'input':'mcq','timed_out':data.timed_out,
          'evaluation':grade_mcq(q.body,transcript)}
    if prior:prior.body=body
    else:db.add(Answer(id=aid,session_id=sid,body=body))
    db.commit()
    return {'id':aid,**body}


@app.patch('/api/answers/{aid}/transcript')
def approve(aid: str, data: TranscriptInput, authorization: str = Header(''), db=Depends(database)):
    a, s = owned(Answer, aid, db, authorization)
    monitored(s)
    duration = a.body['duration']
    from .db import InterviewInteraction
    followup=db.get(InterviewInteraction,a.body['question_id']+':followup')
    if followup and followup.body.get('reply'):
        a.body=a.body|{'followup':{'question':followup.body['prompt'],'reply':followup.body['reply']}}
    a.body = a.body | {'transcript': data.transcript, 'approved': data.approved, 'version': a.body['version']+1,
        'evaluation': {'status': 'pending'}, 'metrics': a.body['metrics'] | {'wpm': round(len(data.transcript.split())/duration*60, 1) if duration and a.body['input'] == 'audio' else None}}
    db.commit()
    return {'id': a.id, **a.body}


@app.post('/api/answers/{aid}/evaluate')
def start_evaluation(aid: str, authorization: str = Header(''), db=Depends(database)):
    a, s = owned(Answer, aid, db, authorization)
    monitored(s)
    if not a.body['approved']:
        raise HTTPException(409, 'Approve the transcript first.')
    for j in db.scalars(select(Job).where(Job.session_id == s.id)):
        if j.body['answer_id'] == aid and j.body['kind'] == 'evaluate' and j.body['version'] == a.body['version'] and j.body['status'] in ['queued', 'running', 'done']:
            return {'id': j.id}
    job = Job(session_id=s.id, body={'kind': 'evaluate', 'answer_id': aid, 'version': a.body['version'], 'status': 'queued'})
    db.add(job)
    db.commit()
    return {'id': job.id}


@app.put('/api/answers/{aid}/feelings')
def feelings(aid: str, data: FeelingsInput, authorization: str = Header(''), db=Depends(database)):
    a, s = owned(Answer, aid, db, authorization)
    db.merge(Feelings(id=aid, session_id=s.id, body=data.model_dump() | {'question_id': a.body['question_id']}))
    db.commit()
    return {'saved': True}


@app.post('/api/sessions/{sid}/observations')
def observation(sid: str, data: ObservationInput, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization)
    s = db.scalar(select(Session).where(Session.id==sid).with_for_update())
    config = s.body.get('proctor_config') or load_config(settings.proctor_config_file)
    if data.type == 'device' and data.label not in {'mobile phone','tablet computer','laptop computer','smartwatch','headphones','second monitor'}:
        raise HTTPException(422, 'A device event must identify a supported device, not a face or person.')
    if data.type == 'device' and (data.score is None or data.score < config['device_confidence']):
        raise HTTPException(422, 'Device confidence is below the configured threshold.')
    if data.question_id:
        q = db.get(Question, data.question_id)
        if not q or q.session_id != sid:
            raise HTTPException(404, 'Question not found.')
    oid = f'{sid}-{data.id}'
    prior = db.get(Observation, oid)
    body = describe(data.model_dump(), s.created) | {'received_at':time.time(), 'limitation': 'Client-reported observations require human review and do not establish cheating.'}
    config = {**config, 'device_warning_only': not validation_status()['passed']}
    if prior:
        body = prior.body | {'duration': max(prior.body['duration'], data.duration)}
    db.merge(Observation(id=oid, session_id=sid, body=body))
    db.flush()
    action=apply_policy(db,s,oid,body,config,prior is None)
    row=db.get(Observation,oid);row.body=row.body|{'action_taken':action,'candidate_id':s.body.get('user_id'),'examination_id':s.body.get('exam_id')}
    if data.type=='gaze_sample':
        activity=dict(s.body.get('proctor_summary') or summarize([],config))
        counts=dict(activity.get('counts',{}));gaze=dict(activity.get('gaze',{}))
        if prior is None:
            counts['gaze_sample']=counts.get('gaze_sample',0)+1
            if data.gaze_direction not in (None,'off_screen'):
                n=gaze.get('samples',0);center=gaze.get('center_samples',round((gaze.get('eye_contact_percent') or 0)*n/100))
                center+=int(data.gaze_direction=='center');n+=1
                gaze={'samples':n,'center_samples':center,'eye_contact_percent':round(100*center/n,1)}
        activity=activity|{'counts':counts,'gaze':gaze}
    else:
        activity = summarize([o.body for o in db.scalars(select(Observation).where(Observation.session_id == sid))], config)
    s.body = s.body | {'proctor_summary':activity}
    if activity['terminate'] and s.mode == 'hr' and not s.body.get('exam_id'):
        s.status = 'terminated'
    db.commit()
    return {'id': oid, 'summary':activity,'security':s.body.get('security',{}),'status':s.status,'server_now':time.time()}


@app.websocket('/api/sessions/{sid}/gaze')
async def gaze_stream(ws:WebSocket,sid:str):
    from .db import LoginSession, User
    from .accounts import blocked
    expected=('https' if ws.url.scheme=='wss' else 'http')+'://'+ws.headers.get('host','')
    if ws.headers.get('origin')!=expected:await ws.close(code=1008);return
    await ws.accept()
    try:
        hello=await asyncio.wait_for(ws.receive_json(),10)
        login_id=hashed(ws.cookies.get('interview_auth',''))
        with DB() as db:
            login=db.get(LoginSession,login_id)
            if not login or login.expires<=time.time() or not hmac.compare_digest(login.csrf,hashed(str(hello.get('csrf','')))):
                await ws.close(code=1008);return
        window_start=time.monotonic();messages=0
        while True:
            raw=await asyncio.wait_for(ws.receive_text(),30)
            if len(raw)>3000:await ws.close(code=1009);return
            data=ObservationInput.model_validate(json.loads(raw))
            if time.monotonic()-window_start>=10:window_start=time.monotonic();messages=0
            messages+=1
            if messages>45:await ws.close(code=1008);return
            if data.type not in ['gaze_sample','gaze_violation']:await ws.close(code=1008);return
            def process_gaze():
                with DB() as db:
                    login=db.get(LoginSession,login_id);user=db.get(User,login.user_id) if login else None
                    if not login or login.expires<=time.time() or not user or blocked(user):raise HTTPException(403,'Session unavailable.')
                    reset=actor.set(user)
                    try:
                        session=active_session(sid,db,'Bearer '+str(hello.get('token','')))
                        if session.mode!='hr' or not session.body.get('exam_started_at'):raise HTTPException(403,'Interview not started.')
                        result=observation(sid,data,'Bearer '+str(hello.get('token','')),db)
                    finally:actor.reset(reset)
                return result
            from starlette.concurrency import run_in_threadpool
            result=await run_in_threadpool(process_gaze)
            await ws.send_json(result)
    except (WebSocketDisconnect,asyncio.TimeoutError,ValueError,HTTPException):
        try:await ws.close(code=1008)
        except (RuntimeError, WebSocketDisconnect):pass


@app.post('/api/sessions/{sid}/face-warning/acknowledge')
def acknowledge_face_warning(sid:str,authorization:str=Header(''),db=Depends(database)):
    s=active_session(sid,db,authorization)
    state=dict(s.body.get('security',{}));state['face_warning_acknowledged']=state.get('face_warning_count',0)
    s.body=s.body|{'security':state};db.commit();return {'status':s.status,'security':state}


@app.post('/api/sessions/{sid}/device-warning/acknowledge')
def acknowledge_device_warning(sid:str,authorization:str=Header(''),db=Depends(database)):
    s=active_session(sid,db,authorization)
    state=dict(s.body.get('security',{}));state['device_warning_acknowledged']=state.get('device_warning_count',0)
    s.body=s.body|{'security':state};db.commit();return {'status':s.status,'security':state}


@app.post('/api/sessions/{sid}/device-warning')
@app.post('/api/interview-sessions/{sid}/device-warning')
def record_device_warning(sid:str,data:dict,authorization:str=Header(''),db=Depends(database)):
    s=active_session(sid,db,authorization)
    s=db.scalar(select(Session).where(Session.id==sid).with_for_update())
    oid=f"{sid}-{data.get('event_id','')}"
    event=db.get(Observation,oid)
    if not event or event.session_id!=sid or event.body.get('type')!='device':raise HTTPException(422,'A confirmed device event is required.')
    device_warning(db,s,oid,warning_only=not validation_status()['passed']);db.commit()
    return {'status':s.status,'security':s.body.get('security',{}),'device_warning_count':s.body.get('device_warning_count',0)}


@app.post('/api/sessions/{sid}/violation')
def terminate_for_violation(sid: str, data: ViolationInput, authorization: str = Header(''), db=Depends(database)):
    """Final one-attempt action for an assigned HR interview only."""
    if data.type in ('device_detected','multiple_faces'):raise HTTPException(409,'Face and device detections use independent warning flows.')
    s=active_session(sid,db,authorization)
    if s.mode!='hr' or not s.body.get('exam_id') or not s.body.get('exam_started_at'):
        raise HTTPException(409,'Violations apply only after an assigned interview begins.')
    state=dict(s.body.get('security',{}))
    state |= {'attempt_used':True,'terminated_reason':data.type,'terminated_at':time.time(),'lock_until':None,'review_required':True}
    s.status='terminated'
    s.body=s.body|{'security':state,'attempt_used':True,'terminated_reason':data.type,'ended_at':time.time()}
    for assignment in db.scalars(select(Assignment).where(Assignment.session_id==s.id)):
        assignment.body=assignment.body|{'status':'terminated','attempt_used':True,'terminated_reason':data.type}
    system_event(db,s,'exam_terminated',reason=data.type,attempt_used=True)
    notify(db,s,'malpractice:'+data.type,'One-attempt interview terminated; review required','high')
    db.commit()
    return {'status':s.status,'security':state}


@app.post('/api/observations/{oid}/snapshot')
async def snapshot(oid: str, file: UploadFile = File(...), authorization: str = Header(''), db=Depends(database)):
    o=db.get(Observation,oid)
    if not o:raise HTTPException(404,'Event not found.')
    # An event can itself terminate the exam. Permit its pending evidence only;
    # this does not restore examination or account access.
    s=active_session(o.session_id,db,authorization,False,evidence_only=True)
    if s.status=='complete':raise HTTPException(409,'Assessment is finalized.')
    if o.body['type'] != 'device' or s.mode != 'hr':
        raise HTTPException(422, 'Only device events accept pictures.')
    raw = await file.read(500*1024+1)
    if len(raw) > 500*1024:
        raise HTTPException(413, 'Picture exceeds 500 KB.')
    try:
        im = Image.open(io.BytesIO(raw))
        if im.format != 'JPEG' or im.width*im.height > 4_000_000:
            raise ValueError()
        im.verify()
        im = Image.open(io.BytesIO(raw)).convert('RGB')
        out = io.BytesIO()
        im.save(out, 'JPEG', quality=80)  # Strip EXIF and embedded metadata.
    except Exception:
        raise HTTPException(422, 'Invalid JPEG image.')
    key = f'{s.id}/{oid}.jpg'
    # Commit key before remote write so failed/partial uploads remain discoverable by cleanup.
    db.merge(Snapshot(id=oid, session_id=s.id, body={'key': key, 'ready': False,'captured_at':time.time()}))
    db.commit()
    try:
        storage.put(key, out.getvalue())
    except Exception:
        raise HTTPException(503, 'Evidence upload failed. Monitoring must pause until retry succeeds.')
    row = db.get(Snapshot, oid)
    row.body = row.body | {'ready': True}
    db.commit()
    return {'id': oid}


@app.patch('/api/observations/{oid}/explanation')
def explanation(oid: str, data: ExplanationInput, authorization: str = Header(''), db=Depends(database)):
    o, s = owned(Observation, oid, db, authorization)
    o.body = o.body | {'explanation': data.explanation}
    db.commit()
    return {'saved': True}


@app.post('/api/sessions/{sid}/complete')
def complete(sid: str, authorization: str = Header(''), db=Depends(database)):
    s = active_session(sid, db, authorization, False)
    if s.status=='terminated':raise HTTPException(409,'Examination terminated.')
    if s.status != 'complete':
        monitored(s)
        answers = db.scalars(select(Answer).where(Answer.session_id == sid)).all()
        question_ids = set(db.scalars(select(Question.id).where(Question.session_id == sid)))
        if not question_ids or len(answers) != len(question_ids) or {a.body.get('question_id') for a in answers} != question_ids or any(not a.body['approved'] or a.body['evaluation']['status'] == 'pending' for a in answers):
            raise HTTPException(409, 'Every question requires an approved answer with a finished evaluation attempt.')
        for o in db.scalars(select(Observation).where(Observation.session_id == sid)):
            if o.body['type'] == 'device' and s.mode == 'hr':
                snap = db.get(Snapshot, o.id)
                if not snap or not snap.body.get('ready'):
                    raise HTTPException(409, 'A device evidence picture has not finished uploading.')
        summary = aggregate([a.body['evaluation'] for a in answers])
        s.body = s.body | {'summary': summary,'ended_at':time.time()}
        s.status = 'complete'
        from .rounds import complete_session
        complete_session(db,s)
        if s.body.get('application_id'):
            from .db import Application
            from .applications import transition
            application=db.get(Application,s.body['application_id'])
            if application and application.status=='interview_scheduled':
                transition(db,application,'interviewed',db.get(User,s.body['hr_id']))
        for role in ['candidate', 'hr'] if s.mode == 'hr' else ['candidate']:
            if not db.scalar(select(Report).where(Report.token_hash==digest(report_token(sid,role)))):
                db.add(Report(session_id=sid, role=role, token_hash=digest(report_token(sid, role)), body={'finalized': time.time()}))
        notify(db,s,'completed','Interview completed; report ready','low')
        db.commit()
    return {role: f'/report/{report_token(sid, role)}' for role in (['candidate', 'hr'] if s.mode == 'hr' else ['candidate'])}


@app.get('/api/reports/{token}')
def report(token: str, db=Depends(database)):
    r, s = authorized_report(token, db)
    questions = sorted(db.scalars(select(Question).where(Question.session_id == s.id)).all(), key=lambda q: q.body['sequence'])
    answers = db.scalars(select(Answer).where(Answer.session_id == s.id)).all()
    summary = aggregate([effective_evaluation(a) for a in answers])
    if r.role == 'hr':
        summary.pop('overall', None)
    observations = [{'id': o.id, **o.body, 'snapshot': bool((snap := db.get(Snapshot, o.id)) and snap.body.get('ready'))}
        for o in db.scalars(select(Observation).where(Observation.session_id == s.id)) if o.body['type'] != 'camera_emotion']
    return {'role': r.role, 'mode': s.mode, 'name': s.body['name'], 'target_role': s.body['role'], 'expires': s.expires,
        'summary': summary, 'questions': [{'id': q.id, **q.body} for q in questions],
        'answers': [{'id': a.id, **a.body, 'ai_evaluation':a.body['evaluation'],'evaluation':effective_evaluation(a)} for a in answers], 'observations': observations,
        'proctor_summary': s.body.get('proctor_summary') or summarize(observations, s.body.get('proctor_config') or load_config(settings.proctor_config_file)),
        'feelings_available': r.role == 'candidate' or s.body['share_feelings'], 'share_feelings': s.body['share_feelings'],
        'monitoring': {k: v for k, v in s.body.get('monitoring', {}).items() if k in ['face_coverage', 'object_coverage']},
        'hr_link': f'/report/{report_token(s.id, "hr")}' if r.role == 'candidate' and s.mode == 'hr' else None}


@app.get('/api/reports/{token}/feelings')
def feelings_report(token: str, db=Depends(database)):
    r, s = authorized_report(token, db)
    if r.role == 'hr' and not s.body['share_feelings']:
        raise HTTPException(403, 'Candidate has not consented to sharing this report.')
    return {'title': 'Candidate self-reported feelings', 'responses': [f.body for f in db.scalars(select(Feelings).where(Feelings.session_id == s.id))]}


@app.patch('/api/reports/{token}/feelings-consent')
def consent(token: str, data: ConsentInput, db=Depends(database)):
    r, s = authorized_report(token, db, True)
    if s.mode != 'hr':
        raise HTTPException(422, 'Practice reports are private to the candidate.')
    s.body = s.body | {'share_feelings': data.share}
    db.add(Consent(session_id=s.id, body={'share': data.share, 'version': '2026-09-v1', 'timestamp': time.time()}))
    db.commit()
    return {'share_feelings': data.share}


@app.get('/api/reports/{token}/snapshots/{snapshot_id}')
def get_snapshot(token: str, snapshot_id: str, db=Depends(database)):
    r, s = authorized_report(token, db)
    snap = db.get(Snapshot, snapshot_id)
    if not snap or snap.session_id != s.id or not snap.body.get('ready'):
        raise HTTPException(404, 'Picture not found.')
    try:
        return Response(storage.get(snap.body['key']), media_type='image/jpeg')
    except Exception:
        raise HTTPException(503, 'Picture temporarily unavailable.')


@app.delete('/api/reports/{token}')
def delete_report(token: str, db=Depends(database)):
    user = require_user('hr', 'admin')
    r, s = authorized_report(token, db)
    if r.role != 'hr' or s.body.get('hr_id') not in organization_recruiters(user):
        raise HTTPException(403, 'Only the assigned HR team can delete this interview.')
    try:
        purge(db, s)
        return {'deleted': True}
    except Exception:
        return {'deleted': False, 'access_revoked': True, 'message': 'Deletion queued for retry.'}


static = Path(__file__).resolve().parents[2] / 'frontend' / 'dist'
if static.exists():
    @app.get('/gaze.js')
    def gaze_script():
        return FileResponse(static / 'gaze.js', media_type='application/javascript',headers={'Cache-Control':'no-cache'})

    @app.get('/device-monitor.js')
    def device_monitor():
        return FileResponse(static / 'device-monitor.js', media_type='application/javascript', headers={'Cache-Control':'no-cache'})

    @app.get('/detector-check')
    def detector_check():
        require_user('hr','admin')
        return FileResponse(static / 'detector-check.html')

    @app.get('/face-monitor.js')
    def face_monitor():
        return FileResponse(static / 'face-monitor.js', media_type='application/javascript')
    app.mount('/assets', StaticFiles(directory=static / 'assets'), name='assets')
    for directory in ['vision', 'models']:
        if (static / directory).exists():
            app.mount('/' + directory, StaticFiles(directory=static / directory), name=directory)
    @app.get('/{path:path}')
    def frontend(path: str):
        if path.startswith('api/'):
            raise HTTPException(404)
        return FileResponse(static / 'index.html')
