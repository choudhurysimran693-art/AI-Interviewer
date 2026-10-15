"""CPU-only resume evidence matching. Findings never change hiring stages."""
import io
import re
import time
from datetime import date
from typing import Literal
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy import select, delete
from .db import DB, RoleCriterion, ResumeExtraction, ScreeningSummary, ScreeningResult, CandidateResume, ApplicationResume, Application, RecruitmentPosition, User
from .accounts import require_user, organization_scope

router = APIRouter(prefix='/api')
MONTHS={name:i+1 for i,name in enumerate(['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec'])}
DATE_TOKEN=r'(?:\d{4}-\d{2}|(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{4})'

def normalized_date(value):
    if value.lower() in ('present','current'):return value.lower()
    if re.fullmatch(r'\d{4}-\d{2}',value):return value
    month,year=value.split();return f'{year}-{MONTHS[month[:3].lower()]:02}'
ALIASES = {'javascript':['javascript','js'], 'typescript':['typescript','ts'], 'python':['python'], 'postgresql':['postgresql','postgres'], 'react':['react','reactjs','react.js'], 'node.js':['node.js','nodejs'], 'machine learning':['machine learning','ml'], 'sql':['sql'], 'java':['java'], 'c++':['c++'], 'aws':['aws','amazon web services']}

def contains(text, value):
    aliases = next((v for v in ALIASES.values() if value.lower() in v), [value.lower()])
    return any(re.search(r'(?<!\w)'+re.escape(a)+r'(?!\w)', text, re.I) for a in aliases)

def source_text(document):
    if document.filename.lower().endswith('.pdf'):
        from pypdf import PdfReader
        return '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(document.content)).pages)
    from docx import Document
    doc=Document(io.BytesIO(document.content))
    return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])

def extract(text):
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    education=[];experience=[];projects=[];certifications=[]
    section=''
    for line in lines:
        heading=line.lower().strip(': ')
        if heading in ('projects','project','education','experience','employment','work experience','professional experience','certifications','certification','skills'):section='experience' if 'experience' in heading else heading
        if re.search(r'\b(b\.?tech|bachelor|master|m\.?tech|b\.?sc|m\.?sc|ph\.?d|diploma)\b',line,re.I):
            years=re.findall(r'\b(?:19|20)\d{2}\b',line)
            education.append({'degree':line,'institution':'Not provided','graduation_year':int(years[-1]) if years else None,'evidence':line})
        dates=re.search('('+DATE_TOKEN+r')\s*(?:to|[-–—])\s*('+DATE_TOKEN+'|present|current)',line,re.I)
        if dates and section=='education':
            entry=next((e for e in reversed(education) if e['evidence']==line),None)
            if entry is None:
                entry={'degree':'Not provided','institution':'Not provided','graduation_year':None,'evidence':line};education.append(entry)
            entry.update(start_date=normalized_date(dates[1]),end_date=normalized_date(dates[2]))
        if dates and section in ('experience','employment'):
            experience.append({'company':'Not provided','role':'Not provided','start_date':normalized_date(dates[1]),'end_date':normalized_date(dates[2]),'description':line,'evidence':line})
        if section.startswith('project') and line.lower() not in ('project','projects'):projects.append(line)
        if re.search(r'\b(certified|certification|certificate)\b',line,re.I) and len(line)>15:certifications.append(line)
    return {'source_text':text,'extracted_skills':[k for k in ALIASES if contains(text,k)],'extracted_education':education,'extracted_experience':experience,'extracted_projects':projects,'extracted_certifications':certifications,'extracted_at':time.time(),'candidate_confirmed':False}

def ensure_extraction(db,document):
    row=db.get(ResumeExtraction,document.id)
    if not row:
        row=ResumeExtraction(id=document.id,candidate_id=document.user_id,body=extract(source_text(document)));db.add(row);db.flush()
    return row

def timeline(entries, today=None):
    today=today or date.today();ranges=[];unclear=False
    for e in entries:
        try:
            start=date.fromisoformat(e['start_date']+'-01')
            end=today if e['end_date'].lower() in ('present','current') else date.fromisoformat(e['end_date']+'-01')
            if end<start or start>today or end>today:raise ValueError()
            ranges.append((start,end,e.get('evidence','')))
        except (ValueError,KeyError,TypeError,AttributeError):unclear=True
    merged=[]
    for start,end,_ in sorted(ranges):
        if merged and start<=merged[-1][1]:merged[-1]=(merged[-1][0],max(end,merged[-1][1]))
        else:merged.append((start,end))
    gaps=[{'start':a[1].isoformat(),'end':b[0].isoformat(),'status':'needs_review','message':'Unexplained interval between documented employment dates; invite candidate context.'} for a,b in zip(merged,merged[1:]) if (b[0]-a[1]).days>31]
    return round(sum((b-a).days for a,b in merged)/365.25,2),gaps,unclear

def compare(criteria, extraction):
    text=extraction.get('source_text','');lines=[x for x in text.splitlines() if x.strip()]
    employment=extraction.get('extracted_experience',[])
    education=extraction.get('extracted_education',[])
    years,_,unclear_dates=timeline(employment)
    _,gaps,_=timeline(employment+[e for e in education if e.get('start_date') and e.get('end_date')])
    for gap in gaps:gap['message']='Unexplained interval between documented employment/education dates; invite candidate context.'
    results=[]
    for c in criteria:
        value=c['criterion_value'];kind=c['criterion_type'];evidence=[];status='not_met';reason='No supporting evidence found in this resume.'
        if kind=='min_experience_years':
            evidence=[e.get('evidence','') for e in extraction.get('extracted_experience',[]) if e.get('evidence') in text and e.get('evidence')]
            status='needs_review';reason='Complete dated employment and role relevance require HR review.'
            if evidence and not unclear_dates:
                status='partially_supported' if years>=float(value) else 'needs_review'
                reason=f'{years} non-overlapping documented years; HR must verify relevance to this role.'
        else:
            pool=lines
            if kind=='qualification':pool=[e.get('evidence','') for e in education]
            elif kind=='certification':pool=extraction.get('extracted_certifications',[])
            evidence=[line for line in pool if line and line in text and contains(line,value)]
            if evidence:
                applied=any(re.search(r'\b(built|developed|implemented|used|designed|managed|created|certified|awarded|degree|bachelor|master)\b',x,re.I) for x in evidence)
                status='matched' if applied or kind in ('qualification','certification') else 'partially_supported'
                reason='Direct source evidence.' if status=='matched' else 'Mention found; depth or use needs confirmation.'
            elif not text:status='needs_review';reason='No readable resume is available.'
        results.append({**c,'status':status,'evidence_text':'\n'.join(evidence[:4]),'reason':reason,'confidence_score':1.0 if evidence else None,'confidence_meaning':'Literal citation lookup certainty; not candidate suitability.'})
    required=[r for r in results if r['criterion_type']=='required_skill']
    strengths=[r['criterion_value'] for r in results if r['status']=='matched']
    missing=[r['criterion_value'] for r in results if r['status']=='not_met' and r['is_mandatory']]
    unclear=[r['criterion_value'] for r in results if r['status'] in ('needs_review','partially_supported')]
    questions=[f"Please describe your experience with {v}, including a concrete example." for v in (missing+unclear+strengths)[:5]]
    for q in ['Which project best demonstrates your fit for this role?','What were your own responsibilities and measurable results?','Which requirements would you need support learning?']:
        if len(questions)<3:questions.append(q)
    return {'results':results,'match_score':None,'strengths_summary':strengths,'missing_summary':missing,'unclear_claims_summary':unclear,'suggested_questions':questions,'documented_experience_years':years if extraction.get('extracted_experience') else None,'required_supported':sum(r['status']=='matched' for r in required),'required_total':len(required),'review_count':len(unclear),'gaps':gaps,'years_since_graduation':[date.today().year-e['graduation_year'] for e in extraction.get('extracted_education',[]) if isinstance(e.get('graduation_year'),int) and e['graduation_year']<=date.today().year],'candidate_confirmed':extraction.get('candidate_confirmed',False),'generated_at':time.time()}

def screen(db,a):
    link=db.get(ApplicationResume,a.id)
    document=db.get(CandidateResume,link.resume_id) if link else None
    extraction=ensure_extraction(db,document).body if document else {}
    criteria=[{'id':c.id,**c.body} for c in db.scalars(select(RoleCriterion).where(RoleCriterion.position_id==a.position_id))]
    previous=db.get(ScreeningSummary,a.id)
    body=compare(criteria,extraction);body['gap_explanation']=previous.body.get('gap_explanation','') if previous else ''
    db.execute(delete(ScreeningResult).where(ScreeningResult.application_id==a.id))
    for r in body['results']:db.add(ScreeningResult(application_id=a.id,criterion_id=r['id'],body=r))
    if previous:previous.body=body
    else:db.add(ScreeningSummary(id=a.id,body=body))
    return body

class CriterionInput(BaseModel):
    criterion_type:Literal['required_skill','preferred_skill','min_experience_years','qualification','certification']
    criterion_value:str=Field(min_length=1,max_length=200)
    is_mandatory:bool=True

class CriteriaInput(BaseModel):
    criteria:list[CriterionInput]=Field(max_length=50)

def position_owned(db,pid,user):
    p=db.get(RecruitmentPosition,pid)
    if not p or p.organization_id!=organization_scope(user):raise HTTPException(404,'Position not found.')
    return p

@router.get('/positions/{pid}/criteria')
def get_criteria(pid:str):
    user=require_user('hr','admin')
    with DB() as db:
        position_owned(db,pid,user)
        return {'criteria':[{'id':c.id,**c.body} for c in db.scalars(select(RoleCriterion).where(RoleCriterion.position_id==pid))]}

@router.post('/positions/{pid}/criteria')
@router.patch('/positions/{pid}/criteria')
def save_criteria(pid:str,data:CriteriaInput):
    user=require_user('hr','admin')
    for c in data.criteria:
        if c.criterion_type=='min_experience_years':
            try:
                if not 0<=float(c.criterion_value)<=60:raise ValueError()
            except ValueError:raise HTTPException(422,'Experience must be a number between 0 and 60.')
    with DB() as db:
        position_owned(db,pid,user)
        apps=list(db.scalars(select(Application).where(Application.position_id==pid)))
        for a in apps:db.execute(delete(ScreeningResult).where(ScreeningResult.application_id==a.id))
        db.execute(delete(RoleCriterion).where(RoleCriterion.position_id==pid))
        for c in data.criteria:db.add(RoleCriterion(position_id=pid,body=c.model_dump()))
        db.flush()
        for a in apps:screen(db,a)
        db.commit();return {'saved':True}

@router.get('/candidate/profile/resume-extraction')
def get_extraction():
    user=require_user('candidate')
    with DB() as db:
        document=db.get(CandidateResume,user.body.get('resume_id',''))
        if not document or document.user_id!=user.id:raise HTTPException(404,'Upload a resume first.')
        row=ensure_extraction(db,document);db.commit();return {'id':row.id,**row.body}

class Correction(BaseModel):
    extracted_skills:list[str]=Field(default_factory=list,max_length=100)
    extracted_education:list[dict]=Field(default_factory=list,max_length=30)
    extracted_experience:list[dict]=Field(default_factory=list,max_length=50)
    extracted_projects:list[str]=Field(default_factory=list,max_length=50)
    extracted_certifications:list[str]=Field(default_factory=list,max_length=50)

@router.patch('/candidate/profile/resume-extraction')
def correct_extraction(data:Correction):
    user=require_user('candidate')
    with DB() as db:
        doc=db.get(CandidateResume,user.body.get('resume_id',''))
        if not doc or doc.user_id!=user.id:raise HTTPException(404,'Upload a resume first.')
        row=ensure_extraction(db,doc)
        for entry in data.extracted_experience+data.extracted_education:
            if not entry.get('evidence') or entry['evidence'] not in row.body['source_text']:raise HTTPException(422,'Each education/employment entry needs an exact source excerpt in evidence.')
        row.body=row.body|data.model_dump()|{'candidate_confirmed':True}
        for a in db.scalars(select(Application).where(Application.candidate_id==user.id)):
            link=db.get(ApplicationResume,a.id)
            if link and link.resume_id==doc.id:screen(db,a)
        db.commit();return {'saved':True}

@router.get('/applications/{aid}/screening')
@router.post('/applications/{aid}/screen')
def screening(aid:str):
    user=require_user('hr','admin')
    from .applications import owned
    with DB() as db:
        a=owned(db,aid,user);body=screen(db,a);db.commit();return body

class InformationRequest(BaseModel):
    message:str=Field(min_length=5,max_length=2000)

@router.post('/applications/{aid}/request-information')
def request_information(aid:str,data:InformationRequest):
    user=require_user('hr','admin')
    from .applications import owned
    from .candidate_notices import add_notice
    from .db import uid
    with DB() as db:
        a=owned(db,aid,user)
        add_notice(db,a.candidate_id,'screening-request:'+uid(),'HR requested resume information',data.message,'/candidate/profile',user.id)
        db.commit();return {'saved':True}

@router.delete('/positions/{pid}/criteria')
def clear_criteria(pid:str):
    return save_criteria(pid,CriteriaInput(criteria=[]))
